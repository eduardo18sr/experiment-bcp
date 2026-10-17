"""Generador de expedientes sintéticos.

Tres casos deliberadamente distintos (framework F2: "se diseña desde dos y se valida con un tercero"):

- CASO_A_PD_CURVA: curva PD temporal con calibración por plazo. Reproduce, con datos inventados, el caso
  controlado real: el código carga una versión de parámetros que no coincide (0/70), la versión que coincide
  está comentada (70/70), la extensión lineal de la curva final solo coincide en los periodos 3-12 y los
  límites de rango difieren (<=19/<=25 frente a 19-24/25-36).
- CASO_B_PRICING: tasa y cuota por segmento con VLOOKUP. El punto de referencia coincide, pero un spread
  del tarifario activo difiere de la referencia: solo lo detecta la comparación de tablas (R6).
- CASO_C_SCORING: scorecard logístico escrito en otra forma algebraica. Todo coincide (control de falsos
  positivos).

Cada caso trae además su documentación metodológica en `documentacion/`: A en Word con estilos en español e
incompleta (el perímetro omite canal y vigencia, y falta la versión de parámetros), B en Markdown completa y
C en Word con estilos en inglés completa.

Los valores en caché del Excel se calculan con una implementación independiente de la especificación, para
que la validación del oráculo contra la caché sea una prueba real y no circular.

Uso: python -m evals.sintetico.generador <directorio_destino>
"""

from __future__ import annotations

import csv
import math
import re
import shutil
import sys
import textwrap
import zipfile
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import yaml

PERIODOS = list(range(3, 73))  # 70 periodos, columnas I..BZ


def _col(numero: int) -> str:
    letras = ""
    while numero > 0:
        numero, resto = divmod(numero - 1, 26)
        letras = chr(65 + resto) + letras
    return letras


COLUMNAS_PERIODO = [_col(9 + i) for i in range(len(PERIODOS))]  # I(9) .. BZ(78)


def _redondear(valor: float, digitos: int) -> float:
    return float(Decimal(repr(valor)).quantize(Decimal(1).scaleb(-digitos), rounding=ROUND_HALF_UP))


def _logistica(x: float) -> float:
    return 1.0 / (1.0 + math.exp(-x))


def _escribir_csv(ruta: Path, encabezados: list[str], filas: list[list]) -> None:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    with open(ruta, "w", newline="", encoding="utf-8") as archivo:
        escritor = csv.writer(archivo)
        escritor.writerow(encabezados)
        for fila in filas:
            escritor.writerow([repr(v) if isinstance(v, float) else v for v in fila])


def _escribir_python(ruta: Path, codigo: str) -> None:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    ruta.write_text(textwrap.dedent(codigo).lstrip(), encoding="utf-8")


def _inyectar_cache(ruta_xlsx: Path, hojas: list[str], cache: dict[str, float]) -> None:
    """openpyxl escribe fórmulas sin valor en caché: se inyecta `<v>` como lo haría Excel al guardar."""
    temporal = ruta_xlsx.with_suffix(".tmp")
    with zipfile.ZipFile(ruta_xlsx) as origen, zipfile.ZipFile(temporal, "w", zipfile.ZIP_DEFLATED) as destino:
        for elemento in origen.infolist():
            contenido = origen.read(elemento.filename)
            coincidencia = re.match(r"xl/worksheets/sheet(\d+)\.xml$", elemento.filename)
            if coincidencia:
                hoja = hojas[int(coincidencia.group(1)) - 1]
                texto = contenido.decode("utf-8")

                def reemplazar(m: re.Match) -> str:
                    ref = f"{hoja}!{m.group(1)}"
                    if ref not in cache:
                        return m.group(0)
                    return f'<c r="{m.group(1)}"{m.group(2)}><f>{m.group(3)}</f><v>{repr(float(cache[ref]))}</v></c>'

                texto = re.sub(r'<c r="([A-Z]+\d+)"([^>]*)><f>(.*?)</f><v\s*/></c>', reemplazar, texto)
                contenido = texto.encode("utf-8")
            destino.writestr(elemento, contenido)
    temporal.replace(ruta_xlsx)


# --------------------------------------------------------------------------------------------------------
# CASO A · curva PD temporal
# --------------------------------------------------------------------------------------------------------
CALIBRACION = [(18.0, 0.10, 1.00), (24.0, 0.00, 1.00), (36.0, 0.00, 1.00)]
RAMPAS_ALPHA = [(-3.0, 0.5), (-2.8, 0.6), (-2.6, 0.7)]  # (alpha en t=3, alpha desde t=18) por rango
DESPLAZAMIENTO_VERSION_ACTIVA = 0.5
FACTOR_FLAT = 2.0


def _theta(desplazamiento: float = 0.0) -> dict[int, dict[str, float]]:
    theta = {}
    for t in PERIODOS:
        fila = {}
        for rango, (inicio, fin) in enumerate(RAMPAS_ALPHA, start=1):
            alpha = inicio + (fin - inicio) * (min(t, 18) - 3) / 15
            fila[f"alpha_{rango}"] = round(alpha + desplazamiento, 10)
            fila[f"beta_{rango}"] = 1.0
        theta[t] = fila
    return theta


def verdad_caso_a(score: float, plazo: int) -> dict:
    """Especificación de la referencia (la calculadora), implementada de forma independiente."""
    factor = 40 / math.log(2)
    offset = 400 - factor * math.log(50)
    xb = -(score - offset) / factor
    rango = 1 if plazo <= 19 else 2 if plazo <= 25 else 3
    _, intercepto, pendiente = CALIBRACION[rango - 1]
    xb_cal = intercepto + pendiente * xb
    theta = _theta()
    curva = [_logistica(theta[t][f"alpha_{rango}"] + theta[t][f"beta_{rango}"] * xb_cal) for t in PERIODOS]
    pd_12 = curva[PERIODOS.index(12)]
    incremento = (FACTOR_FLAT * pd_12 - pd_12) / 24
    final = []
    for i, valor in enumerate(curva):
        final.append(valor if PERIODOS[i] <= 12 else final[-1] + incremento)
    return {
        "offset": offset,
        "factor": factor,
        "xb": xb,
        "pd_base": _logistica(xb),
        "xb_cal": xb_cal,
        "pd_cal": _logistica(xb_cal),
        "score_cal": _redondear(-xb_cal * factor + offset, 0),
        "curva": curva,
        "final": final,
        "objetivo_36": FACTOR_FLAT * pd_12,
        "incremento": incremento,
    }


def _caso_a(destino: Path) -> None:
    import openpyxl

    score, plazo = 300.0, 18
    verdad = verdad_caso_a(score, plazo)
    theta = _theta()

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Calculadora"
    etiquetas = {6: "Offset", 7: "Factor", 8: "Plazo", 9: "xBeta", 10: "PD base", 11: "Score", 12: "xBeta calibrado",
                 13: "PD calibrada", 14: "Score calibrado", 34: "Factor flat"}
    for fila, texto in etiquetas.items():
        hoja[f"G{fila}"] = texto
    hoja["H6"] = "=400-H7*LN(50)"
    hoja["H7"] = "=40/LN(2)"
    hoja["H8"] = plazo
    hoja["H9"] = "=-(H11-H6)/H7"
    hoja["H10"] = "=1/(1+EXP(-H9))"
    hoja["H11"] = score
    hoja["H12"] = "=IF(H8<=19,L9+M9*H9,IF(H8<=25,L10+M10*H9,L11+M11*H9))"
    hoja["H13"] = "=1/(1+EXP(-H12))"
    hoja["H14"] = "=ROUND(-H12*H7+H6,0)"
    for columna, encabezado in zip("KLM", ("plazo_max", "intercepto", "pendiente")):
        hoja[f"{columna}8"] = encabezado
    for indice, fila_cal in enumerate(CALIBRACION):
        for columna, valor in zip("KLM", fila_cal):
            hoja[f"{columna}{9 + indice}"] = valor

    hoja["H20"] = "periodo"
    nombres = ["alpha_1", "beta_1", "alpha_2", "beta_2", "alpha_3", "beta_3"]
    for desplazamiento, nombre in enumerate(nombres):
        hoja[f"H{21 + desplazamiento}"] = nombre
    hoja["H27"] = "pd_temporal"
    hoja["H28"] = "pd_final"
    cache = {
        "Calculadora!H6": verdad["offset"], "Calculadora!H7": verdad["factor"], "Calculadora!H9": verdad["xb"],
        "Calculadora!H10": verdad["pd_base"], "Calculadora!H12": verdad["xb_cal"], "Calculadora!H13": verdad["pd_cal"],
        "Calculadora!H14": verdad["score_cal"], "Calculadora!I34": verdad["objetivo_36"],
        "Calculadora!J34": verdad["incremento"],
    }
    for i, (t, columna) in enumerate(zip(PERIODOS, COLUMNAS_PERIODO)):
        hoja[f"{columna}20"] = t
        for desplazamiento, nombre in enumerate(nombres):
            hoja[f"{columna}{21 + desplazamiento}"] = theta[t][nombre]
        hoja[f"{columna}27"] = (
            f"=IF($H$8<=19,1/(1+EXP(-({columna}21+{columna}22*$H$12))),"
            f"IF($H$8<=25,1/(1+EXP(-({columna}23+{columna}24*$H$12))),1/(1+EXP(-({columna}25+{columna}26*$H$12)))))"
        )
        hoja[f"{columna}28"] = f"={columna}27" if t <= 12 else f"={COLUMNAS_PERIODO[i - 1]}28+$J$34"
        cache[f"Calculadora!{columna}27"] = verdad["curva"][i]
        cache[f"Calculadora!{columna}28"] = verdad["final"][i]
    hoja["H34"] = FACTOR_FLAT
    hoja["I34"] = "=H34*R27"
    hoja["J34"] = "=(I34-R27)/24"

    ruta_excel = destino / "referencia" / "calculadora_pd.xlsx"
    ruta_excel.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta_excel)
    _inyectar_cache(ruta_excel, ["Calculadora"], cache)

    raiz = destino / "implementacion"
    encabezados_theta = ["periodo"] + nombres
    _escribir_csv(raiz / "params" / "theta_Tt_VERSION_EXCEL.csv", encabezados_theta,
                  [[t] + [theta[t][n] for n in nombres] for t in PERIODOS])
    theta_activa = _theta(DESPLAZAMIENTO_VERSION_ACTIVA)
    _escribir_csv(raiz / "params" / "theta_Tt_VERSION_ACTIVA.csv", encabezados_theta,
                  [[t] + [theta_activa[t][n] for n in nombres] for t in PERIODOS])
    _escribir_csv(raiz / "params" / "calibracion_plazo.csv", ["plazo_max", "intercepto", "pendiente"],
                  [list(f) for f in CALIBRACION])
    _escribir_python(raiz / "motor" / "__init__.py", "")
    _escribir_python(raiz / "motor" / "parametros.py", '''
        """Carga de parámetros del motor de PD."""
        import csv
        from pathlib import Path

        RUTA_PARAMS = Path(__file__).resolve().parent.parent / "params"


        def leer_tabla(ruta):
            with open(ruta, newline="", encoding="utf-8") as archivo:
                return [{k: float(v) for k, v in fila.items()} for fila in csv.DictReader(archivo)]


        def cargar_parametros():
            calibracion = leer_tabla(RUTA_PARAMS / "calibracion_plazo.csv")
            # theta = leer_tabla(RUTA_PARAMS / "theta_Tt_VERSION_EXCEL.csv")
            theta = leer_tabla(RUTA_PARAMS / "theta_Tt_VERSION_ACTIVA.csv")
            return calibracion, sorted(theta, key=lambda fila: fila["periodo"])
    ''')
    _escribir_python(raiz / "motor" / "curva.py", '''
        """Curva PD temporal por plazo."""
        import math

        from motor.parametros import cargar_parametros

        FACTOR = 40 / math.log(2)
        OFFSET = 400 - FACTOR * math.log(50)
        FACTOR_FLAT = 2.0


        def _rango_plazo(plazo):
            if plazo <= 18:
                return 1
            if plazo <= 24:
                return 2
            return 3


        def calcular_pd_temporal(score, plazo):
            calibracion, theta = cargar_parametros()
            xb = -(score - OFFSET) / FACTOR
            rango = _rango_plazo(plazo)
            fila = calibracion[rango - 1]
            xb_cal = fila["intercepto"] + fila["pendiente"] * xb
            pd_calibrada = 1 / (1 + math.exp(-xb_cal))
            curva = [
                1 / (1 + math.exp(-(t[f"alpha_{rango}"] + t[f"beta_{rango}"] * xb_cal)))
                for t in theta
            ]
            final = list(curva)
            if plazo > 24:
                pd_12 = curva[9]
                incremento = (FACTOR_FLAT * pd_12 - pd_12) / 24
                for i in range(10, len(final)):
                    final[i] = final[i - 1] + incremento
            return {"pd_calibrada": pd_calibrada, "pd_temporal": curva, "pd_final": final}
    ''')
    expediente = {
        "id": "CASO_A_PD_CURVA",
        "modelo_id": "MOD-PD-001",
        "descripcion": "Curva PD temporal con calibración por plazo (datos sintéticos)",
        "referencia": {"excel": "referencia/calculadora_pd.xlsx"},
        "implementacion": {"raiz": "implementacion", "permitir_ejecucion": True},
        "entradas": [
            {"nombre": "score", "hoja": "Calculadora", "celda": "H11", "tipo": "numero", "paso": 1},
            {"nombre": "plazo", "hoja": "Calculadora", "celda": "H8", "tipo": "entero"},
        ],
        "componentes": [
            {"id": "pd_calibrada", "tipo": "salida", "descripcion": "PD calibrada por plazo",
             "referencia": {"hoja": "Calculadora", "rango": "H13"},
             "implementacion": {"funcion": "motor.curva:calcular_pd_temporal", "salida": "pd_calibrada"}},
            {"id": "curva_pd_temporal", "tipo": "salida", "descripcion": "Curva PD temporal, periodos 3 a 72",
             "referencia": {"hoja": "Calculadora", "rango": "I27:BZ27", "etiquetas": "I20:BZ20"},
             "implementacion": {"funcion": "motor.curva:calcular_pd_temporal", "salida": "pd_temporal"}},
            {"id": "curva_pd_final", "tipo": "derivado", "descripcion": "Extensión lineal desde el periodo 13",
             "referencia": {"hoja": "Calculadora", "rango": "I28:BZ28", "etiquetas": "I20:BZ20"},
             "implementacion": {"funcion": "motor.curva:calcular_pd_temporal", "salida": "pd_final"}},
        ],
        "tablas_parametros": [
            {"id": "theta_temporal", "referencia": {"hoja": "Calculadora", "rango": "H20:BZ26"},
             "orientacion": "filas", "clave": "periodo", "archivos": "params/*.csv"},
            {"id": "calibracion_plazo", "referencia": {"hoja": "Calculadora", "rango": "K8:M11"},
             "orientacion": "columnas", "clave": "plazo_max", "archivos": "params/*.csv"},
        ],
    }
    (destino / "expediente.yaml").write_text(yaml.safe_dump(expediente, sort_keys=False, allow_unicode=True), encoding="utf-8")


# --------------------------------------------------------------------------------------------------------
# CASO B · pricing por segmento
# --------------------------------------------------------------------------------------------------------
TARIFARIO = {
    "MINORISTA": (0.12, 0.035, 0.010),
    "PYME": (0.14, 0.030, 0.015),
    "CORPORATIVO": (0.09, 0.020, 0.005),
    "EMPRESA": (0.11, 0.025, 0.008),
}


def verdad_caso_b(segmento: str, monto: float, plazo: int) -> dict:
    tasa_base, spread, ajuste = TARIFARIO[segmento]
    ajuste = ajuste if plazo > 12 else 0.0
    tasa_anual = tasa_base + spread + ajuste
    tasa_mensual = (1 + tasa_anual) ** (1 / 12) - 1
    cuota = _redondear(monto * tasa_mensual / (1 - (1 + tasa_mensual) ** (-plazo)), 2)
    return {"tasa_base": tasa_base, "spread": spread, "ajuste": ajuste, "tasa_anual": tasa_anual,
            "tasa_mensual": tasa_mensual, "cuota": cuota}


def _caso_b(destino: Path) -> None:
    import openpyxl

    segmento, monto, plazo = "MINORISTA", 10000.0, 24
    verdad = verdad_caso_b(segmento, monto, plazo)
    libro = openpyxl.Workbook()
    calculo = libro.active
    calculo.title = "Calculo"
    tarifario = libro.create_sheet("Tarifario")
    for columna, encabezado in zip("ABCD", ("segmento", "tasa_base", "spread", "ajuste_plazo")):
        tarifario[f"{columna}1"] = encabezado
    for fila, (nombre, valores) in enumerate(TARIFARIO.items(), start=2):
        tarifario[f"A{fila}"] = nombre
        for columna, valor in zip("BCD", valores):
            tarifario[f"{columna}{fila}"] = valor
    for fila, texto in {2: "Segmento", 3: "Monto", 4: "Plazo", 6: "Tasa base", 7: "Spread", 8: "Ajuste plazo",
                        9: "Tasa anual", 10: "Tasa mensual", 11: "Cuota"}.items():
        calculo[f"A{fila}"] = texto
    calculo["B2"], calculo["B3"], calculo["B4"] = segmento, monto, plazo
    calculo["B6"] = "=VLOOKUP(B2,Tarifario!$A$2:$D$5,2,FALSE)"
    calculo["B7"] = "=VLOOKUP(B2,Tarifario!$A$2:$D$5,3,FALSE)"
    calculo["B8"] = "=IF(B4>12,VLOOKUP(B2,Tarifario!$A$2:$D$5,4,FALSE),0)"
    calculo["B9"] = "=B6+B7+B8"
    calculo["B10"] = "=(1+B9)^(1/12)-1"
    calculo["B11"] = "=ROUND(B3*B10/(1-(1+B10)^(-B4)),2)"
    ruta_excel = destino / "referencia" / "tarifario_pricing.xlsx"
    ruta_excel.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta_excel)
    _inyectar_cache(ruta_excel, ["Calculo", "Tarifario"], {
        "Calculo!B6": verdad["tasa_base"], "Calculo!B7": verdad["spread"], "Calculo!B8": verdad["ajuste"],
        "Calculo!B9": verdad["tasa_anual"], "Calculo!B10": verdad["tasa_mensual"], "Calculo!B11": verdad["cuota"],
    })

    raiz = destino / "implementacion"
    encabezados = ["segmento", "tasa_base", "spread", "ajuste_plazo"]
    vigente = {k: list(v) for k, v in TARIFARIO.items()}
    vigente["PYME"][1] = 0.031  # spread divergente en el tarifario que carga el código
    _escribir_csv(raiz / "params" / "tarifario_2026.csv", encabezados, [[k] + v for k, v in vigente.items()])
    _escribir_csv(raiz / "params" / "tarifario_2025.csv", encabezados,
                  [[k, round(v[0] + 0.01, 6), v[1], v[2]] for k, v in TARIFARIO.items()])
    _escribir_python(raiz / "pricing" / "__init__.py", "")
    _escribir_python(raiz / "pricing" / "tarifario.py", '''
        import csv
        from pathlib import Path

        CARPETA = Path(__file__).resolve().parent.parent / "params"


        def cargar_tarifario():
            with open(CARPETA / "tarifario_2026.csv", newline="", encoding="utf-8") as archivo:
                return {
                    fila["segmento"].upper(): {k: float(v) for k, v in fila.items() if k != "segmento"}
                    for fila in csv.DictReader(archivo)
                }
    ''')
    _escribir_python(raiz / "pricing" / "cuota.py", '''
        from decimal import ROUND_HALF_UP, Decimal

        from pricing.tarifario import cargar_tarifario


        def _redondear(valor, digitos):
            return float(Decimal(repr(valor)).quantize(Decimal(1).scaleb(-digitos), rounding=ROUND_HALF_UP))


        def calcular_cuota(segmento, monto, plazo):
            tarifa = cargar_tarifario()[segmento.upper()]
            ajuste = tarifa["ajuste_plazo"] if plazo >= 13 else 0.0
            tasa_anual = tarifa["tasa_base"] + tarifa["spread"] + ajuste
            tasa_mensual = (1 + tasa_anual) ** (1 / 12) - 1
            cuota = monto * tasa_mensual / (1 - (1 + tasa_mensual) ** (-plazo))
            return {"tasa_anual": tasa_anual, "cuota": _redondear(cuota, 2)}
    ''')
    expediente = {
        "id": "CASO_B_PRICING",
        "modelo_id": "MOD-PRC-001",
        "descripcion": "Tasa y cuota por segmento (datos sintéticos)",
        "referencia": {"excel": "referencia/tarifario_pricing.xlsx"},
        "implementacion": {"raiz": "implementacion", "permitir_ejecucion": True},
        "entradas": [
            {"nombre": "segmento", "hoja": "Calculo", "celda": "B2", "tipo": "texto"},
            {"nombre": "monto", "hoja": "Calculo", "celda": "B3", "tipo": "numero", "paso": 1},
            {"nombre": "plazo", "hoja": "Calculo", "celda": "B4", "tipo": "entero"},
        ],
        "componentes": [
            {"id": "tasa_anual", "tipo": "salida", "referencia": {"hoja": "Calculo", "rango": "B9"},
             "implementacion": {"funcion": "pricing.cuota:calcular_cuota", "salida": "tasa_anual"}},
            {"id": "cuota", "tipo": "salida", "referencia": {"hoja": "Calculo", "rango": "B11"},
             "implementacion": {"funcion": "pricing.cuota:calcular_cuota", "salida": "cuota"}},
        ],
        "tablas_parametros": [
            {"id": "tarifario", "referencia": {"hoja": "Tarifario", "rango": "A1:D5"},
             "orientacion": "columnas", "clave": "segmento", "archivos": "params/*.csv"},
        ],
    }
    (destino / "expediente.yaml").write_text(yaml.safe_dump(expediente, sort_keys=False, allow_unicode=True), encoding="utf-8")


# --------------------------------------------------------------------------------------------------------
# CASO C · scorecard logístico
# --------------------------------------------------------------------------------------------------------
COEFICIENTES = {"intercepto": -1.2, "edad": -0.02, "ingreso_log": -0.15, "mora_12m": 0.9}
ENTRADAS_C = {"edad": 35.0, "ingreso_log": 8.2, "mora_12m": 1.0}


def verdad_caso_c() -> dict:
    xb = COEFICIENTES["intercepto"] + sum(COEFICIENTES[k] * v for k, v in ENTRADAS_C.items())
    return {"xbeta": xb, "pd": _logistica(xb)}


def _caso_c(destino: Path) -> None:
    import openpyxl

    verdad = verdad_caso_c()
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Score"
    hoja["A1"], hoja["B1"] = "variable", "coeficiente"
    for fila, (nombre, valor) in enumerate(COEFICIENTES.items(), start=2):
        hoja[f"A{fila}"], hoja[f"B{fila}"] = nombre, valor
    for fila, (nombre, valor) in enumerate(ENTRADAS_C.items(), start=2):
        hoja[f"D{fila}"], hoja[f"E{fila}"] = nombre, valor
    hoja["D6"], hoja["E6"] = "xbeta", "=B2+B3*E2+B4*E3+B5*E4"
    hoja["D7"], hoja["E7"] = "pd", "=1/(1+EXP(-E6))"
    ruta_excel = destino / "referencia" / "scorecard.xlsx"
    ruta_excel.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta_excel)
    _inyectar_cache(ruta_excel, ["Score"], {"Score!E6": verdad["xbeta"], "Score!E7": verdad["pd"]})

    raiz = destino / "implementacion"
    _escribir_csv(raiz / "params" / "coeficientes.csv", ["variable", "coeficiente"], [[k, v] for k, v in COEFICIENTES.items()])
    _escribir_python(raiz / "scoring" / "__init__.py", "")
    _escribir_python(raiz / "scoring" / "modelo.py", '''
        import csv
        import math
        from pathlib import Path

        COEFICIENTES = Path(__file__).resolve().parent.parent / "params" / "coeficientes.csv"


        def _coeficientes():
            with open(COEFICIENTES, newline="", encoding="utf-8") as archivo:
                return {fila["variable"]: float(fila["coeficiente"]) for fila in csv.DictReader(archivo)}


        def calcular_pd(edad, ingreso_log, mora_12m):
            beta = _coeficientes()
            variables = {"edad": edad, "ingreso_log": ingreso_log, "mora_12m": mora_12m}
            xbeta = beta["intercepto"] + sum(beta[nombre] * valor for nombre, valor in variables.items())
            # Forma algebraica equivalente a 1 / (1 + exp(-xbeta)).
            pd = math.exp(xbeta) / (1 + math.exp(xbeta))
            return {"xbeta": xbeta, "pd": pd}
    ''')
    expediente = {
        "id": "CASO_C_SCORING",
        "modelo_id": "MOD-PD-002",
        "descripcion": "Scorecard logístico de admisión (datos sintéticos)",
        "referencia": {"excel": "referencia/scorecard.xlsx"},
        "implementacion": {"raiz": "implementacion", "permitir_ejecucion": True},
        "entradas": [
            {"nombre": "edad", "hoja": "Score", "celda": "E2", "tipo": "numero", "paso": 1},
            {"nombre": "ingreso_log", "hoja": "Score", "celda": "E3", "tipo": "numero"},
            {"nombre": "mora_12m", "hoja": "Score", "celda": "E4", "tipo": "entero"},
        ],
        "componentes": [
            {"id": "xbeta", "tipo": "salida", "referencia": {"hoja": "Score", "rango": "E6"},
             "implementacion": {"funcion": "scoring.modelo:calcular_pd", "salida": "xbeta"}},
            {"id": "pd", "tipo": "salida", "referencia": {"hoja": "Score", "rango": "E7"},
             "implementacion": {"funcion": "scoring.modelo:calcular_pd", "salida": "pd"}},
        ],
        "tablas_parametros": [
            {"id": "coeficientes", "referencia": {"hoja": "Score", "rango": "A1:B5"},
             "orientacion": "columnas", "clave": "variable", "archivos": "params/*.csv"},
        ],
    }
    (destino / "expediente.yaml").write_text(yaml.safe_dump(expediente, sort_keys=False, allow_unicode=True), encoding="utf-8")


# --------------------------------------------------------------------------------------------------------
# Documentación metodológica (dimensión documentación)
# --------------------------------------------------------------------------------------------------------
def _documentacion_a(destino: Path) -> None:
    """Word en español, incompleto: el perímetro omite canal y vigencia, y no documenta la versión de parámetros."""
    from evals.sintetico.docx_minimo import ESTILOS_ESPANOL, escribir_docx

    escribir_docx(destino / "documentacion" / "metodologia_pd.docx", [
        ("titulo", 0, "Documento metodológico · PD temporal de consumo (sintético)"),
        ("titulo", 1, "1. Objetivo del modelo"),
        ("parrafo", "Estimar la probabilidad de incumplimiento por periodo a partir del score de originación y del plazo."),
        ("titulo", 1, "2. Perímetro de aplicación"),
        ("parrafo", "El modelo se aplica a los créditos de consumo del segmento personas naturales, "
                    "producto préstamo personal en cuotas."),
        ("titulo", 1, "3. Variables del modelo"),
        ("tabla", [["Variable", "Descripción", "Fuente"],
                   ["score", "Puntaje de originación del cliente", "Motor de decisión"],
                   ["plazo", "Plazo del crédito en meses", "Solicitud"]]),
        ("titulo", 1, "4. Calibración por plazo"),
        ("parrafo", "La calibración se aplica por rangos de plazo con un intercepto y una pendiente por rango."),
    ], estilos=ESTILOS_ESPANOL)


def _documentacion_b(destino: Path) -> None:
    """Markdown completo."""
    ruta = destino / "documentacion" / "metodologia_pricing.md"
    ruta.parent.mkdir(parents=True, exist_ok=True)
    ruta.write_text(textwrap.dedent("""
        # Metodología de pricing de créditos (sintético)

        ## Perímetro de aplicación

        Segmentos: minorista, PYME, corporativo y empresa. Producto: crédito en cuotas a tasa fija.
        Canal: agencias y banca digital. Vigencia: desde enero de 2026 hasta la siguiente revisión anual.

        ## Tarifario vigente

        El tarifario vigente es la versión 2026, aprobada por el comité de precios y cargada desde el archivo
        de parámetros versionado del proyecto.

        ## Segmentos

        | Segmento | Tasa base | Spread | Ajuste por plazo |
        |---|---|---|---|
        | MINORISTA | 12 % | 3,5 % | 1,0 % |
        | PYME | 14 % | 3,0 % | 1,5 % |
        | CORPORATIVO | 9 % | 2,0 % | 0,5 % |
        | EMPRESA | 11 % | 2,5 % | 0,8 % |
    """).lstrip(), encoding="utf-8")


def _documentacion_c(destino: Path) -> None:
    """Word con estilos en inglés, completo."""
    from evals.sintetico.docx_minimo import ESTILOS_INGLES, escribir_docx

    escribir_docx(destino / "documentacion" / "metodologia_scoring.docx", [
        ("titulo", 0, "Scorecard de admisión (sintético)"),
        ("titulo", 1, "Ámbito de aplicación"),
        ("parrafo", "Segmento: personas naturales sin historial previo en la entidad. Producto: tarjeta de crédito "
                    "clásica. Canal: banca digital. Vigencia: 2026, con monitoreo trimestral."),
        ("titulo", 1, "Versión de parámetros"),
        ("parrafo", "Los coeficientes vigentes están en params/coeficientes.csv, versión aprobada en enero de 2026."),
        ("titulo", 1, "Diccionario de variables"),
        ("tabla", [["Variable", "Descripción"],
                   ["edad", "Edad del solicitante en años"],
                   ["ingreso_log", "Logaritmo del ingreso mensual declarado"],
                   ["mora_12m", "Indicador de mora en los últimos 12 meses"]]),
    ], estilos=ESTILOS_INGLES)


DOCUMENTACION = {"CASO_A_PD_CURVA": _documentacion_a, "CASO_B_PRICING": _documentacion_b, "CASO_C_SCORING": _documentacion_c}


INVENTARIO = {
    "modelos": [
        {"id": "MOD-PD-001", "nombre": "PD temporal de consumo (sintético)", "tipo_modelo": "pd", "tier": 1,
         "responsable": "equipo-modelos-sintetico"},
        {"id": "MOD-PRC-001", "nombre": "Pricing de créditos (sintético)", "tipo_modelo": "pricing", "tier": 2,
         "responsable": "equipo-modelos-sintetico"},
        {"id": "MOD-PD-002", "nombre": "Scoring de admisión (sintético)", "tipo_modelo": "pd", "tier": 3,
         "responsable": "equipo-modelos-sintetico"},
    ]
}

CASOS = {"CASO_A_PD_CURVA": _caso_a, "CASO_B_PRICING": _caso_b, "CASO_C_SCORING": _caso_c}


def generar_casos(destino: Path, casos: list[str] | None = None) -> dict[str, Path]:
    """Genera los expedientes y el inventario en `destino`. Devuelve {id_caso: carpeta}."""
    destino.mkdir(parents=True, exist_ok=True)
    (destino / "inventario.yaml").write_text(yaml.safe_dump(INVENTARIO, sort_keys=False, allow_unicode=True), encoding="utf-8")
    generados = {}
    for caso, constructor in CASOS.items():
        if casos and caso not in casos:
            continue
        carpeta = destino / caso
        if carpeta.exists():
            shutil.rmtree(carpeta)
        carpeta.mkdir(parents=True)
        constructor(carpeta)
        DOCUMENTACION[caso](carpeta)
        generados[caso] = carpeta
    return generados


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(2)
    for caso, carpeta in generar_casos(Path(sys.argv[1])).items():
        print(f"{caso}: {carpeta}")
