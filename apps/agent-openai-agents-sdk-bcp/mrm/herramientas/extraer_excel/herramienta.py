"""Herramienta `extraer_excel`: libro Excel -> LibroCanonico."""

from __future__ import annotations

from pathlib import Path

from pydantic import BaseModel, Field

from mrm.contratos.canonico import CeldaCanonica, LibroCanonico
from mrm.contratos.herramienta import Herramienta, TipoHerramienta
from mrm.herramientas._comun.celdas import sha256_archivo
from mrm.herramientas.extraer_excel.formulas import ErrorFormula, analizar_formula


class EntradaExtraerExcel(BaseModel):
    ruta: str
    hojas: list[str] = Field(default_factory=list)


class SalidaExtraerExcel(BaseModel):
    libro: LibroCanonico


def _texto_formula(valor) -> str | None:
    """openpyxl devuelve str para fórmulas normales y ArrayFormula para matriciales."""
    if isinstance(valor, str) and valor.startswith("="):
        return valor
    texto = getattr(valor, "text", None)
    if isinstance(texto, str):
        return texto if texto.startswith("=") else f"={texto}"
    return None


class ExtraerExcel(Herramienta[EntradaExtraerExcel, SalidaExtraerExcel]):
    nombre = "extraer_excel"
    version = "1.0.0"
    tipo = TipoHerramienta.DETERMINISTICO
    errores = ("archivo_no_encontrado", "formato_no_soportado")
    Entrada = EntradaExtraerExcel
    Salida = SalidaExtraerExcel

    def ejecutar(self, entrada: EntradaExtraerExcel) -> SalidaExtraerExcel:
        import openpyxl

        ruta = Path(entrada.ruta)
        if not ruta.is_file():
            raise self.fallar("archivo_no_encontrado", f"No existe {ruta.name}")
        if ruta.suffix.lower() not in (".xlsx", ".xlsm"):
            raise self.fallar("formato_no_soportado", f"Extensión {ruta.suffix} no soportada")

        # El libro se abre dos veces: con fórmulas y con los valores que Excel dejó en caché.
        libro_formulas = openpyxl.load_workbook(ruta, data_only=False)
        libro_valores = openpyxl.load_workbook(ruta, data_only=True)

        hojas = entrada.hojas or libro_formulas.sheetnames
        celdas: dict[str, CeldaCanonica] = {}
        advertencias: list[str] = []
        formulas_con_cache = 0
        total_formulas = 0

        for hoja in hojas:
            if hoja not in libro_formulas.sheetnames:
                advertencias.append(f"Hoja no encontrada: {hoja}")
                continue
            hoja_f = libro_formulas[hoja]
            hoja_v = libro_valores[hoja]
            for fila in hoja_f.iter_rows():
                for celda in fila:
                    if celda.value is None:
                        continue
                    ref = f"{hoja}!{celda.coordinate}"
                    formula = _texto_formula(celda.value)
                    if formula is None:
                        celdas[ref] = CeldaCanonica(ref=ref, valor=celda.value)
                        continue
                    total_formulas += 1
                    cache = hoja_v[celda.coordinate].value
                    if cache is not None:
                        formulas_con_cache += 1
                    try:
                        arbol = analizar_formula(formula, hoja)
                    except ErrorFormula as error:
                        arbol = None
                        advertencias.append(f"{ref}: fórmula no traducible ({error.codigo})")
                    celdas[ref] = CeldaCanonica(ref=ref, valor=cache, formula=formula, arbol=arbol)

        nombres = {}
        for nombre, definicion in libro_formulas.defined_names.items():
            texto = getattr(definicion, "attr_text", None)
            if texto:
                nombres[nombre] = texto.replace("$", "")

        if total_formulas and formulas_con_cache == 0:
            advertencias.append(
                "El libro no tiene valores en caché (no fue guardado por Excel): el oráculo no puede validarse"
            )

        return SalidaExtraerExcel(
            libro=LibroCanonico(
                archivo=ruta.name,
                sha256=sha256_archivo(ruta),
                hojas=list(hojas),
                celdas=celdas,
                nombres_definidos=nombres,
                tiene_cache=formulas_con_cache > 0,
                advertencias=advertencias,
            )
        )
